# Fills the downloaded import template the way a person would in Excel (openpyxl saves compressed, with real dates/times).
import sys, datetime, warnings
warnings.simplefilter("error")
import openpyxl
src, dst = sys.argv[1], sys.argv[2]
wb = openpyxl.load_workbook(src); ws = wb.active
head = [c.value for c in ws[3]]
col = {h: i + 1 for i, h in enumerate(head)}
ws.cell(1, 2, "Bulk Project")
rows = [
    {"Client": "Ελληνική ΑΕ", "Shipping container": "MSCU1234566", "Date": datetime.date(2026, 10, 27), "Time": datetime.time(9, 30), "Weight (kg)": 21000, "Pieces": 840, "Price (€)": 120.5, "Status": "Scheduled", "Cargo type": "Resin", "Sender (Αποστολέας)": "Αποστολέας ΑΕ", "Receiver (Παραλήπτης)": "Port of Piraeus"},
    {"Client": "Second", "HCT container": "HCTU0000011", "Date": "28/10/2026", "Truck plate": "ikx1234", "Project": "Own project", "Invoice no": "INV-7"},
    {},
    {"Client": "Third", "Weight (kg)": "lots", "Status": "Maybe"},
]
r = 4
for row in rows:
    for k, v in row.items():
        ws.cell(r, col[k], v)
    r += 1
wb.save(dst)
